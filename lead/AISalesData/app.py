"""
app.py  —  AI Sales Machine dashboard  (v4)

v4 CHANGES:
- New "📋 List Campaign" page: upload CSV/Excel with first_name, last_name,
  email (and optionally company) — sends to every row, rotating sender accounts.
- Scraper campaign page unchanged (Google Maps mode still works).
- Both modes share the same Email Composer, sender accounts, and dashboard.
- Placeholders: {first_name}, {last_name}, {full_name}, {company_name}.
- CSV template download button so users know the exact format.
"""

import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import threading
import time
import io
import uuid

import smtplib
import ssl
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from email.utils import formatdate, make_msgid

import database
import send_worker
import campaign_runner
from ai_writer import revise_email
from smtp_sender import (
    SMTPSender, test_sender, PROVIDER_PRESETS, MAX_EMAILS_PER_SENDER,
)
from csv_importer import (
    parse_contact_list, validate_contact_list, generate_sample_csv,
    parse_sender_list, generate_sender_sample_csv, senders_to_csv,
)
from database import (
    insert_lead, insert_email_sent,
    get_campaign_stats, get_all_emails_sent, create_campaign,
    update_campaign_status, save_paused_campaign, get_paused_campaigns,
    save_sender_accounts, load_sender_accounts,
    count_sender_accounts, clear_sender_accounts,
    create_scheduled_send, get_scheduled_sends, cancel_scheduled_send,
    get_replies, count_unread_replies, set_reply_flags,
    save_setting, load_settings, get_daily_report_data,
)
from reply_fetcher import fetch_all_replies, imap_host_for
from datetime import datetime, date, time as dtime, timedelta
import crypto_util
from followup_scheduler import get_followup_candidates, send_followups
from utils import (
    get_recent_logs, clear_log_buffer, log_info, sanitize_credential,
)

def _save_app_setting(key: str) -> None:
    val = st.session_state.get(key, "")
    save_setting(key, str(val))


def _restore_settings_from_db():
    """Load settings from SQLite database and populate st.session_state if not present."""
    try:
        saved = load_settings()
        for k in ("offering", "cname", "cweb", "api_key"):
            if k not in st.session_state and k in saved and saved[k]:
                st.session_state[k] = saved[k]
        if "fu_days" not in st.session_state and "fu_days" in saved:
            try:
                st.session_state["fu_days"] = int(saved["fu_days"])
            except Exception:
                pass
    except Exception:
        pass


# ── PAGE CONFIG ────────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="AI Sales Machine",
    page_icon="📩",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── AUTHENTICATION ─────────────────────────────────────────────────────────────
# Must run before anything else renders: every page below exposes live SMTP
# credentials, the lead database, and the ability to send mail.
import auth
auth.require_login()

# Auto-restore saved settings from SQLite after login passes
_restore_settings_from_db()

st.markdown("""
<style>
.metric-card{
    background:linear-gradient(135deg,#1e2130,#252d3d);
    border:1px solid #333;border-radius:12px;
    padding:20px;text-align:center;margin:4px;
}
.metric-value{font-size:2.2rem;font-weight:bold;color:#00d4aa;}
.metric-label{font-size:.85rem;color:#aaa;margin-top:4px;}
.log-box{
    background:#0d1117;border:1px solid #30363d;border-radius:8px;
    padding:12px;font-family:monospace;font-size:11px;
    height:280px;overflow-y:auto;color:#7ee787;white-space:pre-wrap;
}
.mode-badge-ai{display:inline-block;background:#0d4f3c;color:#00d4aa;
    border-radius:6px;padding:4px 12px;font-size:13px;font-weight:bold;}
.mode-badge-manual{display:inline-block;background:#3d2d00;color:#ffa500;
    border-radius:6px;padding:4px 12px;font-size:13px;font-weight:bold;}
.mode-badge-fallback{display:inline-block;background:#2d1f3d;color:#b39ddb;
    border-radius:6px;padding:4px 12px;font-size:13px;font-weight:bold;}
.list-info{background:#0d2137;border:1px solid #1a4a7a;border-radius:8px;
    padding:14px 18px;margin:10px 0;}
</style>
""", unsafe_allow_html=True)

# ── SECRETS ────────────────────────────────────────────────────────────────────

def get_secret(name: str, default: str = "") -> str:
    """
    Read a value from .streamlit/secrets.toml.

    st.secrets raises StreamlitSecretNotFoundError when no secrets file
    exists anywhere — even from .get() — so every read must be guarded or
    the app crashes for users who never created one.
    """
    try:
        return str(st.secrets[name]).strip()
    except Exception:
        return default


# ── SESSION STATE ──────────────────────────────────────────────────────────────

# Widget-key prefixes used by the per-sender editor. Keys are suffixed with a
# sender's stable uid (never its list index), so deleting or replacing senders
# can't make one row inherit another's widget state.
_SENDER_WIDGET_PREFIXES = (
    "email", "pass", "smtp", "port", "ssl",
    "name", "desig", "scomp", "sweb", "prov",
)


def _new_sender(company_name: str = "", company_website: str = "",
                provider: str = "Gmail") -> dict:
    """Build a blank sender account with a stable uid."""
    preset = PROVIDER_PRESETS.get(provider, PROVIDER_PRESETS["Gmail"])
    return {
        "uid":             uuid.uuid4().hex[:8],
        "email":           "",
        "password":        "",
        "smtp_server":     preset["smtp_server"],
        "smtp_port":       preset["smtp_port"],
        "use_ssl":         preset["use_ssl"],
        "name":            "",
        "designation":     "",
        "company_name":    company_name,
        "company_website": company_website,
        "provider":        provider,
    }


_DEFAULTS = {
    # Campaign progress now lives in campaign_runner's thread-safe state, not
    # here — session_state is unreadable from the background thread.
    "_stop_event":       threading.Event(),   # used by the follow-up send
    "custom_subject":    "",
    "custom_body":       "",
    "uploaded_contacts": [],   # parsed contacts from CSV upload
    "campaign_consumed": False,  # composed emails moved into review yet?
    "sender_test_results": {},  # uid -> {"ok": bool, "error": str}
    "sender_filter":       "",
    "followup_candidates": None,   # None = not previewed; [] = previewed, none found
    "followup_selected":   {},     # email_id -> bool
    "followup_result":     None,   # summary dict after a send
    "campaign_prepared":   None,   # composed-but-unsent emails awaiting approval
    "campaign_selected":   {},     # prepared id -> bool
    "campaign_meta":       {},     # {campaign_id, source}
    "campaign_result":     None,   # summary dict after a send
}

for k, v in _DEFAULTS.items():
    if k not in st.session_state:
        st.session_state[k] = v

# Sender accounts are restored from encrypted storage once per session, so a
# browser refresh no longer means retyping every SMTP password.
if "senders" not in st.session_state:
    try:
        _saved = load_sender_accounts()
    except Exception as exc:
        _saved = []
        st.warning(f"Could not load saved sender accounts: {exc}")
    st.session_state.senders        = _saved or [_new_sender()]
    st.session_state.senders_loaded = bool(_saved)

# Poll for scheduled sends whose time has come. Idempotent per process.
send_worker.start_worker()


def _purge_sender_widgets(uids) -> None:
    """
    Drop the widget state belonging to the given sender uids.

    Streamlit ignores a widget's value= argument once its key exists in
    session_state, so stale keys must be removed or a reused/new sender in
    that position would display the old account's data.
    """
    for uid in uids:
        for prefix in _SENDER_WIDGET_PREFIXES:
            st.session_state.pop(f"{prefix}_{uid}", None)


def _auto_save_senders() -> None:
    """Auto-save sender accounts with details to encrypted SQLite storage."""
    try:
        senders = st.session_state.get("senders", [])
        active = [s for s in senders if s.get("email") or s.get("password")]
        if active:
            save_sender_accounts(senders)
    except Exception:
        pass


def _delete_sender(uid: str) -> None:
    """Remove one sender account and everything keyed to it."""
    st.session_state.senders = [
        s for s in st.session_state.senders if s.get("uid") != uid
    ]
    st.session_state.sender_test_results.pop(uid, None)
    _purge_sender_widgets([uid])
    # Never leave the form with zero rows to edit.
    if not st.session_state.senders:
        st.session_state.senders = [_new_sender(
            st.session_state.get("cname", ""), st.session_state.get("cweb", "")
        )]
    _auto_save_senders()


def _replace_senders(new_senders: list[dict]) -> None:
    """Swap the whole sender list, clearing widget and test state first."""
    _purge_sender_widgets([s.get("uid") for s in st.session_state.senders])
    st.session_state.sender_test_results = {}
    st.session_state.senders = new_senders
    _auto_save_senders()

# ── SIDEBAR ────────────────────────────────────────────────────────────────────

with st.sidebar:
    st.markdown("## AI Sales Machine")
    st.markdown("---")

    # Unread count acts as the admin's notification that someone replied.
    try:
        _unread = count_unread_replies()
    except Exception:
        _unread = 0
    _inbox_label = f"Replies ({_unread})" if _unread else "Replies"

    page = st.radio(
        "=>Navigation",
        [
            "Scraper Campaign",
            "List Campaign",
            _inbox_label,
            "Dashboard",
            "Manage Senders",
        ],
        key="nav_page",
    )
    if _unread:
        st.success(f"✉️ **{_unread} unread reply(ies)** waiting in Replies")

    st.markdown("---")

    _restore_settings_from_db()

    with st.expander("🏢 Company & Offer Setup", expanded=False):
        offering = st.text_area(
            "What do you offer?", key="offering", height=80,
            value=st.session_state.get("offering", ""),
            placeholder="e.g. social media marketing for local businesses",
            on_change=_save_app_setting, args=("offering",)
        )
        company_name = st.text_input(
            "Company Name", key="cname",
            value=st.session_state.get("cname", ""),
            placeholder="GrowthLab Agency",
            on_change=_save_app_setting, args=("cname",)
        )
        company_website = st.text_input(
            "Company Website", key="cweb",
            value=st.session_state.get("cweb", ""),
            placeholder="https://growthlab.com",
            on_change=_save_app_setting, args=("cweb",)
        )

    with st.expander("AI Key & follow-up Settings", expanded=False):
        _secret_api_key = get_secret("openai_api_key") or get_secret("gemini_api_key")

        if _secret_api_key:
            st.success("🔑 API key loaded from `secrets.toml`")
            _typed_key = st.text_input(
                "API Key", type="password", key="api_key",
                value=st.session_state.get("api_key", ""),
                placeholder="Enter API",
                on_change=_save_app_setting, args=("api_key",)
            )
            ai_api_key = _typed_key.strip() or _secret_api_key
        else:
            ai_api_key = st.text_input(
                "API Key", type="password", key="api_key",
                value=st.session_state.get("api_key", ""),
                placeholder="Enter API",
                on_change=_save_app_setting, args=("api_key",)
            ).strip()

        follow_up_days = st.number_input(
            "Follow-up After (days)", 0, 14,
            value=int(st.session_state.get("fu_days", 3)),
            key="fu_days",
            on_change=_save_app_setting, args=("fu_days",)
        )

    # Immediately auto-persist any typed sidebar settings to SQLite DB
    if offering:        save_setting("offering", offering)
    if company_name:    save_setting("cname", company_name)
    if company_website: save_setting("cweb", company_website)
    if ai_api_key:      save_setting("api_key", ai_api_key)
    if follow_up_days is not None: save_setting("fu_days", str(follow_up_days))

    st.markdown("---")
    auth.render_logout(st.sidebar)


# ── EMAIL MODE HELPER ─────────────────────────────────────────────────────────

def _email_mode():
    key = (ai_api_key or "").strip()
    has_api = bool(key and len(key) >= 10)
    has_manual = bool(
        st.session_state.get("custom_subject", "").strip() and
        st.session_state.get("custom_body", "").strip()
    )
    if has_manual: return "manual"
    if has_api:    return "ai"
    return "fallback"


# ── EMAIL COMPOSER ─────────────────────────────────────────────────────────────

def render_email_composer(context="scraper"):
    """
    context = "scraper" : placeholders available are {company_name}
    context = "list"    : placeholders available are {first_name}, {last_name},
                          {full_name}, {company_name}
    """
    mode = _email_mode()
    st.markdown("## ✍️ Email Content")

    badge_map = {
        "ai":       '<span class="mode-badge-ai">🤖 AI Mode — AI (Gemini / OpenAI) writes each email</span>',
        "manual":   '<span class="mode-badge-manual">✏️ Manual Template — your email will be sent</span>',
        "fallback": '<span class="mode-badge-fallback">📋 Fallback — a generic template will be used</span>',
    }
    st.markdown(badge_map[mode], unsafe_allow_html=True)
    st.markdown("")

    if mode == "fallback":
        st.info(
            "💡 **No AI API key provided.**\n\n"
            "Enter your own Subject and Body below *(recommended)*, "
            "or leave blank to send a generic template."
        )

    if mode == "ai":
        with st.expander("✏️ Override with your own template (optional)"):
            _render_template_fields(context)
    else:
        _render_template_fields(context)


def _render_template_fields(context):
    if context == "list":
        st.markdown(
            "**Placeholders you can use:**  "
            "`{first_name}` `{last_name}` `{full_name}` `{company_name}`  \n"
            "Each one is replaced with the contact's actual data from your CSV."
        )
    else:
        st.markdown(
            "**Placeholder:** `{company_name}` is replaced with each prospect's business name."
        )

    new_subject = st.text_input(
        "📌 Email Subject",
        value=st.session_state.get("custom_subject", ""),
        key="subject_input",
        placeholder="e.g.  Quick question for {first_name}" if context == "list"
                    else "e.g.  Quick question for {company_name}",
    )

    default_placeholder = (
        "Hi {first_name},\n\n"
        "I came across your details and thought we could help with [your offer].\n\n"
        "Would you be open to a quick 10-minute call this week?\n\n"
        "Best regards,\n[Your Name]"
    ) if context == "list" else (
        "Hi {company_name} team,\n\n"
        "I came across your business and thought we could help with [your offer].\n\n"
        "Would you be open to a quick 10-minute call?\n\n"
        "Best regards,\n[Your Name]"
    )

    new_body = st.text_area(
        "📝 Email Body",
        value=st.session_state.get("custom_body", ""),
        key="body_input",
        height=200,
        placeholder=default_placeholder,
    )

    # Live preview
    if new_subject or new_body:
        with st.expander("👁️ Preview", expanded=False):
            preview_sub  = new_subject.replace("{first_name}", "John") \
                                      .replace("{last_name}", "Smith") \
                                      .replace("{full_name}", "John Smith") \
                                      .replace("{company_name}", "Acme Corp")
            preview_body = new_body.replace("{first_name}", "John") \
                                   .replace("{last_name}", "Smith") \
                                   .replace("{full_name}", "John Smith") \
                                   .replace("{company_name}", "Acme Corp")
            st.markdown(f"**Subject:** {preview_sub or '*(empty)*'}")
            st.markdown("---")
            st.text(preview_body or "*(empty)*")

    st.session_state["custom_subject"] = new_subject
    st.session_state["custom_body"]    = new_body


# ── SENDER FORM ────────────────────────────────────────────────────────────────

def _render_sender_summary(senders):
    """Read-only overview of every configured sender. Never shows passwords."""
    results = st.session_state.sender_test_results
    rows = []
    for i, s in enumerate(senders, 1):
        r = results.get(s.get("uid"))
        if   r is None:  status = "—"
        elif r["ok"]:    status = "✅ OK"
        else:            status = "❌ Failed"
        rows.append({
            "#":          i,
            "Email":      s.get("email") or "(not set)",
            "Provider":   s.get("provider", "Custom / Other"),
            "Name":       s.get("name", ""),
            "Ready":      "✅" if (s.get("email") and s.get("password")) else "⚠️",
            "Last test":  status,
        })
    st.dataframe(pd.DataFrame(rows), use_container_width=True,
                 hide_index=True, height=min(360, 40 + 35 * len(rows)))

    ready = sum(1 for s in senders if s.get("email") and s.get("password"))
    c1, c2, c3 = st.columns(3)
    c1.metric("📧 Accounts", len(senders))
    c2.metric("✅ Ready to send", ready)
    c3.metric("📬 Session capacity", f"{ready * MAX_EMAILS_PER_SENDER:,}")
    if ready < len(senders):
        st.caption(
            f"⚠️ {len(senders) - ready} account(s) are missing an email or "
            "password and will be skipped when a campaign runs."
        )


def _test_all_senders(senders) -> None:
    """Test every configured sender, recording results against their uid."""
    active = [s for s in senders if s.get("email") and s.get("password")]
    if not active:
        st.warning("No accounts have both an email and a password yet.")
        return

    bar  = st.progress(0.0)
    note = st.empty()
    ok = failed = 0
    for n, s in enumerate(active, 1):
        note.caption(f"Testing {s['email']} ({n}/{len(active)})…")
        try:
            r = test_sender(s)
        except Exception as exc:
            # One unexpected failure must not abort the rest of the batch.
            r = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
        st.session_state.sender_test_results[s["uid"]] = {
            "ok": r["ok"], "error": r.get("error", "")
        }
        ok, failed = (ok + 1, failed) if r["ok"] else (ok, failed + 1)
        bar.progress(n / len(active))
    note.empty()
    bar.empty()
    st.success(f"Tested {len(active)} account(s) — ✅ {ok} passed, ❌ {failed} failed.")


def _render_sender_storage(senders):
    """Save / restore / forget encrypted sender accounts on this machine."""
    try:
        saved_count = count_sender_accounts()
    except Exception as exc:
        st.error(f"Saved-account storage unavailable: {exc}")
        return

    with_creds = [s for s in senders if s.get("email") and s.get("password")]

    with st.expander(
        f"💾 Saved accounts on this computer ({saved_count} stored)",
        expanded=False,
    ):
        st.markdown(
            "Accounts are stored in the local database with their passwords "
            "**encrypted**, and restored automatically next time you open "
            "the app — so a refresh no longer means retyping every password."
        )

        try:
            if crypto_util.uses_local_keyfile():
                st.warning(
                    "⚠️ The encryption key is stored **next to the database** "
                    "on this machine. That protects the database file if it "
                    "is copied or backed up on its own, but anyone who can "
                    "read both files can decrypt. For real separation, put a "
                    "key in `.streamlit/secrets.toml` as "
                    "`sender_encryption_key` — see the setup note below."
                )
            else:
                st.success(f"🔐 Encryption key source: {crypto_util.key_source()}")
        except ValueError as exc:
            st.error(f"❌ {exc}")
            return

        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("💾 Save accounts", use_container_width=True,
                         type="primary", key="save_senders"):
                if not with_creds:
                    st.warning("No account has both an email and a password yet.")
                else:
                    try:
                        n = save_sender_accounts(senders)
                        st.session_state.sender_import_done = (
                            f"Saved {n} account(s) to this computer (encrypted)."
                        )
                        st.rerun()
                    except Exception as exc:
                        st.error(f"Could not save: {exc}")
        with c2:
            if st.button("♻️ Restore saved", use_container_width=True,
                         disabled=saved_count == 0, key="restore_senders"):
                restored = load_sender_accounts()
                if restored:
                    _replace_senders(restored)
                    st.session_state.sender_import_done = (
                        f"Restored {len(restored)} saved account(s)."
                    )
                    st.rerun()
                else:
                    st.warning("Nothing saved to restore.")
        with c3:
            if st.button("🗑️ Forget saved", use_container_width=True,
                         disabled=saved_count == 0, key="forget_senders"):
                n = clear_sender_accounts()
                st.session_state.sender_import_done = (
                    f"Deleted {n} saved account(s) from this computer. "
                    "The accounts still in the form above are untouched."
                )
                st.rerun()

        st.caption(
            f"{len(with_creds)} of {len(senders)} account(s) in the form have "
            "a password and would be saved. Saving mirrors the list above — "
            "accounts you removed here are removed from storage too."
        )

        st.markdown("---")
        st.markdown("**Using your own encryption key (recommended)**")
        st.markdown(
            "Create `.streamlit/secrets.toml` next to `app.py`:\n\n"
            "```toml\n"
            'openai_api_key = "sk-..."\n'
            'sender_encryption_key = "<paste generated key>"\n'
            "```\n\n"
            "Generate the key with:\n\n"
            "```bash\n"
            "python -c \"from cryptography.fernet import Fernet; "
            "print(Fernet.generate_key().decode())\"\n"
            "```"
        )
        st.caption(
            "⚠️ Changing the key later makes already-saved passwords "
            "unreadable — the accounts remain listed, but you will need to "
            "re-enter and re-save their passwords."
        )


def _render_sender_import():
    """CSV/Excel bulk import, template download, and export."""
    with st.expander("📥 Import sender accounts from CSV / Excel", expanded=False):
        st.markdown(
            "Bulk-add sending accounts. Only **email** and **password** are "
            "required — the SMTP server is filled in from the `provider` "
            "column, or inferred from the email domain for common providers."
        )
        st.caption(
            "🔒 A sender file contains plain-text SMTP passwords. Keep it "
            "somewhere safe and delete it when the import is done."
        )
        st.download_button(
            "⬇️ Download sender CSV template",
            data=generate_sender_sample_csv(),
            file_name="senders_template.csv",
            mime="text/csv",
            key="dl_sender_template",
        )

        uploaded = st.file_uploader(
            "Upload sender list",
            type=["csv", "xlsx", "xls"],
            key="sender_file",   # must differ from the contacts uploader
            help="Columns: email, password (required); provider, smtp_server, "
                 "smtp_port, use_ssl, name, designation, company, website",
        )

        if uploaded is None:
            return

        parsed, report = parse_sender_list(uploaded.read(), uploaded.name)
        if not parsed:
            st.error(
                "❌ No usable sender accounts found. Every row needs a valid "
                "`email` and a `password`, plus a resolvable SMTP server."
            )
            if report["total"]:
                st.caption(
                    f"Read {report['total']} row(s): "
                    f"{report['invalid_email']} invalid email, "
                    f"{report['no_password']} missing password, "
                    f"{report['no_server']} unresolved server, "
                    f"{report['duplicate']} duplicate."
                )
            return

        st.success(f"✅ Parsed **{report['imported']}** account(s) from **{uploaded.name}**")
        skipped = (report["invalid_email"] + report["no_password"]
                   + report["no_server"] + report["duplicate"])
        if skipped:
            st.warning(
                f"⚠️ Skipped {skipped} row(s): "
                f"{report['invalid_email']} invalid email, "
                f"{report['no_password']} missing password, "
                f"{report['no_server']} could not resolve an SMTP server, "
                f"{report['duplicate']} duplicate address."
            )

        preview = pd.DataFrame([
            {"Email": s["email"], "Provider": s["provider"],
             "SMTP": f"{s['smtp_server']}:{s['smtp_port']}",
             "SSL": "Yes" if s["use_ssl"] else "No", "Name": s["name"]}
            for s in parsed[:10]
        ])
        st.dataframe(preview, use_container_width=True, hide_index=True)
        if len(parsed) > 10:
            st.caption(f"… and {len(parsed) - 10} more not shown.")

        mode = st.radio(
            "How should these be applied?",
            ["Append to my existing accounts", "Replace all existing accounts"],
            key="sender_import_mode",
            horizontal=True,
        )
        test_after = st.checkbox(
            "Test every connection after importing",
            value=False, key="sender_import_test",
            help="Performs a real SMTP login per account — slow for a large file.",
        )

        if not st.button("✅ Apply import", key="apply_sender_import",
                         type="primary"):
            return

        if mode.startswith("Replace"):
            for s in parsed:
                s["uid"] = uuid.uuid4().hex[:8]
            _replace_senders(parsed)
            st.session_state.sender_import_done = (
                f"Replaced all sender accounts with {len(parsed)} imported account(s)."
            )
        else:
            existing = st.session_state.senders
            # Drop the untouched blank starter row so it doesn't linger.
            if len(existing) == 1 and not existing[0].get("email"):
                _purge_sender_widgets([existing[0].get("uid")])
                existing = []
            have = {s.get("email", "").lower() for s in existing if s.get("email")}
            added, dupes = [], 0
            for s in parsed:
                if s["email"] in have:
                    dupes += 1
                    continue
                have.add(s["email"])
                s["uid"] = uuid.uuid4().hex[:8]
                added.append(s)
            st.session_state.senders = existing + added
            msg = f"Added {len(added)} new sender account(s)."
            if dupes:
                msg += f" Skipped {dupes} already in your list."
            st.session_state.sender_import_done = msg

        if test_after:
            _test_all_senders(st.session_state.senders)
        st.rerun()


def _render_sender_editors(senders, preset_names):
    visible = list(enumerate(senders, 1))
    if len(senders) > 3:
        needle = st.text_input(
            "🔎 Filter accounts",
            key="sender_filter",
            placeholder="Search by email, name, or company…",
        ).strip().lower()

        if needle:
            visible = [
                (i, s) for i, s in visible
                if needle in " ".join([
                    s.get("email", ""), s.get("name", ""),
                    s.get("company_name", ""), s.get("provider", ""),
                ]).lower()
            ]
            st.caption(f"Showing {len(visible)} of {len(senders)} accounts.")
            if not visible:
                st.info("No accounts match that search.")

    expand_first = len(senders) <= 3
    for i, sender in visible:
        uid    = sender["uid"]
        label  = f"📧 Sender #{i}: {sender['email'] or 'Not configured'}"
        result = st.session_state.sender_test_results.get(uid)
        if result:
            label += "  ✅" if result["ok"] else "  ❌"

        with st.expander(label, expanded=(expand_first and i == 1)):
            current_provider = sender.get("provider", "Custom / Other")
            chosen_provider  = st.selectbox(
                "Email Provider", preset_names,
                index=preset_names.index(current_provider)
                      if current_provider in preset_names else 0,
                key=f"prov_{uid}",
            )
            if chosen_provider != current_provider:
                p = PROVIDER_PRESETS[chosen_provider]
                sender.update({"provider": chosen_provider,
                               "smtp_server": p["smtp_server"],
                               "smtp_port": p["smtp_port"],
                               "use_ssl": p["use_ssl"]})
                for pref in ("smtp", "port", "ssl"):
                    st.session_state.pop(f"{pref}_{uid}", None)
                st.rerun()

            note_text = PROVIDER_PRESETS.get(chosen_provider, {}).get("note", "")
            if note_text:
                st.caption(f"ℹ️ {note_text}")

            c1, c2 = st.columns(2)
            with c1:
                sender["email"]       = st.text_input("Email Address",    value=sender["email"],          key=f"email_{uid}", placeholder="you@domain.com")
                sender["smtp_server"] = st.text_input("SMTP Server",      value=sender["smtp_server"],    key=f"smtp_{uid}",  placeholder="smtp.gmail.com")
                sender["smtp_port"]   = st.number_input("SMTP Port",      value=int(sender["smtp_port"]), key=f"port_{uid}", min_value=1, max_value=65535)
                sender["use_ssl"]     = st.checkbox("Use SSL (port 465)", value=bool(sender.get("use_ssl", False)), key=f"ssl_{uid}")
            with c2:
                sender["password"]       = sanitize_credential(st.text_input("Password / App Password", value=sender["password"], key=f"pass_{uid}", type="password", placeholder="App password"))
                sender["name"]           = st.text_input("Your Name",               value=sender["name"],        key=f"name_{uid}",  placeholder="John Smith")
                sender["designation"]    = st.text_input("Job Title",               value=sender["designation"], key=f"desig_{uid}", placeholder="Sales Manager")
                sender["company_name"]   = st.text_input("Company (signature)",     value=sender.get("company_name", company_name),       key=f"scomp_{uid}", placeholder="Your Company")
                sender["company_website"]= st.text_input("Website (signature)",     value=sender.get("company_website", company_website), key=f"sweb_{uid}",  placeholder="https://yoursite.com")

            if result and not result["ok"]:
                st.error(f"❌ Last test failed: {result['error']}")

            b1, b2 = st.columns([1, 1])
            with b1:
                if st.button("🔌 Test Connection", key=f"test_{uid}",
                             use_container_width=True):
                    with st.spinner("Testing…"):
                        r = test_sender(sender)
                    st.session_state.sender_test_results[uid] = {
                        "ok": r["ok"], "error": r.get("error", "")
                    }
                    st.rerun()
            with b2:
                if st.button("🗑️ Remove", key=f"del_{uid}",
                             use_container_width=True):
                    _delete_sender(uid)
                    st.rerun()


def render_sender_form(allow_bulk: bool = False):
    st.markdown("## ⚙️ Email Sender Accounts")
    st.info("💡 Works with Gmail, Hostinger, cPanel, Outlook, Zoho, and more. "
            "Emails rotate across all configured accounts automatically.")

    done = st.session_state.pop("sender_import_done", None)
    if done:
        st.success(f"✅ {done}")

    senders      = st.session_state.senders
    preset_names = list(PROVIDER_PRESETS.keys())

    ready = sum(1 for s in senders if s.get("email") and s.get("password"))
    c1, c2, c3 = st.columns(3)
    c1.metric("📧 Accounts", len(senders))
    c2.metric("✅ Ready to send", ready)
    c3.metric("📬 Session capacity", f"{ready * MAX_EMAILS_PER_SENDER:,}")
    if ready < len(senders):
        st.caption(
            f"⚠️ {len(senders) - ready} account(s) are missing an email or "
            "password and will be skipped when a campaign runs."
        )

    st.markdown("<br>", unsafe_allow_html=True)

    if allow_bulk:
        tab1, tab2, tab3, tab4 = st.tabs([
            "⚙️ Manage Accounts",
            "📋 Accounts Overview",
            "💾 Save & Backup",
            "📥 Import & Export",
        ])

        with tab1:
            ca, cb = st.columns(2)
            with ca:
                if st.button("🔌 Test All Connections", use_container_width=True, type="primary"):
                    _test_all_senders(senders)
            with cb:
                if st.button("➕ Add New Sender", use_container_width=True):
                    senders.append(_new_sender(company_name, company_website))
                    st.session_state.senders = senders
                    st.rerun()

            st.markdown("---")
            _render_sender_editors(senders, preset_names)

        with tab2:
            _render_sender_summary(senders)

        with tab3:
            _render_sender_storage(senders)

        with tab4:
            st.download_button(
                "⬇️ Export accounts (no passwords)",
                data=senders_to_csv(senders),
                file_name="senders_export.csv",
                mime="text/csv",
                use_container_width=True,
                help="Exports server settings and names. Passwords are NOT included.",
            )
            st.markdown("---")
            _render_sender_import()
    else:
        _render_sender_editors(senders, preset_names)
        st.markdown("")
        if st.button("➕ Add Sender", use_container_width=True):
            senders.append(_new_sender(company_name, company_website))
            st.rerun()

    st.session_state.senders = senders
    _auto_save_senders()


# ── SHARED SEND LOOP ───────────────────────────────────────────────────────────

def _launch_campaign(source, **params):
    """
    Hand a campaign to the background runner.

    Nothing here blocks: the thread survives Streamlit reruns, so an
    incidental click during a long scrape can no longer destroy the run.
    """
    _auto_save_senders()
    params["source"] = source
    params["senders"] = st.session_state.senders
    if not campaign_runner.start_campaign(params):
        st.warning("A campaign is already running.")
        return
    st.session_state.campaign_prepared = None
    st.session_state.campaign_selected = {}
    st.session_state.campaign_consumed = False


def _adopt_prepared_for_review(snap):
    """
    Move composed emails from runner state into session state once, so the
    review screen can edit them without the runner owning the edits.
    """
    if st.session_state.get("campaign_consumed"):
        return
    prepared = snap["prepared"]
    batch = uuid.uuid4().hex[:6]
    for p in prepared:
        p["id"] = f"{batch}:{p['id']}"
    st.session_state["cp_versions"]     = {}
    st.session_state.campaign_prepared  = prepared
    st.session_state.campaign_selected  = {p["id"]: True for p in prepared}
    st.session_state.campaign_meta      = {"campaign_id": snap["campaign_id"],
                                           "source": snap["source"]}
    st.session_state.campaign_consumed  = True



# ── SHARED LAUNCH CONTROLS ─────────────────────────────────────────────────────

def _render_email_review(items, prefix, selected, extra_columns=None):
    """
    Shared approval screen: show the exact text going to each recipient,
    allow editing it by hand or via an AI instruction, and let each
    recipient be ticked or skipped.

    Mutates *items* in place as the user edits, and *selected*
    (id -> bool) as they tick. Returns the list of chosen items; the
    caller owns the actual send.
    """
    # ── Bulk selection ─────────────────────────────────────────────────────
    c1, c2, c3 = st.columns([1, 1, 3])
    with c1:
        if st.button("Select all", use_container_width=True, key=f"{prefix}_all"):
            for it in items:
                selected[it["id"]] = True
            st.rerun()
    with c2:
        if st.button("Select none", use_container_width=True, key=f"{prefix}_none"):
            for it in items:
                selected[it["id"]] = False
            st.rerun()

    # ── Recipient table ────────────────────────────────────────────────────
    rows = []
    for it in items:
        row = {
            "Send":      bool(selected.get(it["id"], True)),
            "Recipient": it["recipient"],
            "Subject":   it["subject"],
        }
        for label, key in (extra_columns or {}).items():
            row[label] = it.get(key, "")
        rows.append(row)

    edited = st.data_editor(
        pd.DataFrame(rows),
        use_container_width=True,
        hide_index=True,
        height=min(360, 80 + 35 * len(items)),
        key=f"{prefix}_editor",
        column_config={
            "Send": st.column_config.CheckboxColumn(
                "Send", help="Untick to skip this recipient", default=True
            ),
        },
        disabled=[c for c in rows[0].keys() if c != "Send"] if rows else [],
    )
    for idx, it in enumerate(items):
        selected[it["id"]] = bool(edited.iloc[idx]["Send"])

    # ── Per-recipient editor ───────────────────────────────────────────────
    st.markdown("#### ✏️ Review and edit")
    # The label must not depend on the subject: editing a subject would change
    # the option text and Streamlit could no longer match its stored selection.
    pick = st.selectbox(
        "Showing the email for:",
        range(len(items)),
        format_func=lambda i: f'{i + 1}. {items[i]["recipient"]}',
        key=f"{prefix}_pick",
    )
    item = items[pick]
    uid  = item["id"]

    # Once a user has touched a widget, Streamlit ignores value= on later
    # reruns and keeps serving what the browser holds — deleting the
    # session_state key does not reset the live widget. So an AI revision is
    # published by bumping this counter, which changes the widget key and
    # therefore creates a genuinely new widget that honours value=.
    versions = st.session_state.setdefault(f"{prefix}_versions", {})
    ver = versions.get(uid, 0)

    new_subject = st.text_input(
        "Subject", value=item["subject"], key=f"{prefix}_subj_{uid}_v{ver}"
    )
    new_body = st.text_area(
        "Body", value=item["body"], height=230, key=f"{prefix}_body_{uid}_v{ver}"
    )
    # Hand edits are authoritative — write them straight back.
    item["subject"] = new_subject
    item["body"]    = new_body

    st.caption(
        "Your signature is appended automatically when the email is sent, "
        "so there is no need to add one here."
    )

    # ── AI revision ────────────────────────────────────────────────────────
    with st.expander("🤖 Ask AI to revise this email", expanded=False):
        instruction = st.text_input(
            "What should change?",
            key=f"{prefix}_instr_{uid}",
            placeholder="e.g. make it shorter and less formal",
        )
        a1, a2 = st.columns(2)
        with a1:
            revise_one = st.button(
                "✨ Revise this email", use_container_width=True,
                key=f"{prefix}_revise_{uid}",
            )
        with a2:
            revise_all = st.button(
                "✨ Apply to all selected", use_container_width=True,
                key=f"{prefix}_revise_all_{uid}",
                help="Runs the same instruction over every ticked recipient "
                     "— one AI call each.",
            )

        if revise_one or revise_all:
            targets = ([it for it in items if selected.get(it["id"])]
                       if revise_all else [item])
            if not targets:
                st.warning("Nothing is selected.")
            else:
                bar, note = st.progress(0.0), st.empty()
                failures = []
                for n, target in enumerate(targets, 1):
                    note.caption(f"Revising {n}/{len(targets)} — {target['recipient']}")
                    result, err = revise_email(
                        api_key=ai_api_key,
                        subject=target["subject"],
                        body=target["body"],
                        instruction=instruction,
                        recipient_hint=target.get("display_name", ""),
                    )
                    if result:
                        target["subject"] = result["subject"]
                        target["body"]    = result["body"]
                        # New widget identity on the next run, so the revised
                        # text actually reaches the screen.
                        versions[target["id"]] = versions.get(target["id"], 0) + 1
                    else:
                        failures.append(f"{target['recipient']}: {err}")
                    bar.progress(n / len(targets))
                note.empty(); bar.empty()

                if failures:
                    for f in failures[:3]:
                        st.error(f"❌ {f}")
                if len(failures) < len(targets):
                    st.session_state[f"{prefix}_revised_msg"] = (
                        f"AI revised {len(targets) - len(failures)} email(s)."
                    )
                    st.rerun()

    msg = st.session_state.pop(f"{prefix}_revised_msg", None)
    if msg:
        st.success(f"✅ {msg}")

    return [it for it in items if selected.get(it["id"])]


def _render_campaign_review():
    """
    Approval screen for a campaign: the composed emails are shown and can be
    edited before anything is sent. Reached after the contacts are gathered
    and the copy is generated, never before.
    """
    prepared = st.session_state.campaign_prepared
    if prepared is None:
        return

    st.markdown("---")
    st.markdown("### 📧 Review before sending")

    if not prepared:
        st.warning("No emails could be composed for these contacts.")
        if st.button("Close", key="cp_close_empty"):
            st.session_state.campaign_prepared = None
            st.rerun()
        return

    meta = st.session_state.campaign_meta or {}
    st.info(
        f"**{len(prepared)} email(s)** are composed and ready. "
        "**Nothing has been sent yet.** Edit any message below — by hand or "
        "by asking AI — untick anyone you want to skip, then confirm."
    )

    chosen = _render_email_review(
        prepared, prefix="cp",
        selected=st.session_state.campaign_selected,
        extra_columns={"Company": "company_name"},
    )

    st.markdown("---")
    active = [s for s in st.session_state.senders if s.get("email") and s.get("password")]
    capacity = len(active) * MAX_EMAILS_PER_SENDER
    if len(chosen) > capacity:
        st.warning(
            f"⚠️ {len(chosen)} selected but your {len(active)} sender(s) can "
            f"only send {capacity} this session — the remainder would fail."
        )

    c1, c2 = st.columns([2, 1])
    with c1:
        send_clicked = st.button(
            f"📨 Send {len(chosen)} email(s) now",
            type="primary",
            disabled=not chosen or not active,
            use_container_width=True,
            key="cp_send",
        )
    with c2:
        if st.button("Cancel campaign", use_container_width=True, key="cp_cancel"):
            cid = meta.get("campaign_id")
            if cid:
                try: update_campaign_status(cid, "cancelled")
                except Exception: pass
            st.session_state.campaign_prepared = None
            st.session_state.campaign_selected = {}
            st.session_state.campaign_meta     = {}
            st.rerun()

    if send_clicked:
        # Sending also runs on the background thread, so a long send survives
        # reruns and STOP stays clickable throughout.
        campaign_runner.start_send(
            chosen, st.session_state.senders, meta.get("campaign_id")
        )
        st.session_state.campaign_prepared = None
        st.session_state.campaign_selected = {}
        st.session_state.campaign_meta     = {}
        st.session_state.campaign_consumed = False
        st.rerun()


def _render_followup_review():
    """
    Review-then-send panel for follow-ups.

    Clicking "Run Follow-ups" only builds a preview; nothing is sent until
    the user ticks recipients here and confirms. Previously the button
    mailed every eligible contact immediately with no count and no undo.
    """
    candidates = st.session_state.followup_candidates
    if candidates is None:
        return

    st.markdown("---")
    st.markdown("### 🔄 Follow-up Review")

    if not candidates:
        st.success(
            "✅ No follow-ups are due right now. Nothing has been sent."
        )
        if st.button("Close", key="fu_close_empty"):
            st.session_state.followup_candidates = None
            st.rerun()
        return

    selected = st.session_state.followup_selected

    st.info(
        f"**{len(candidates)} recipient(s)** are due a follow-up. "
        "Nothing has been sent yet — review and edit the text below, untick "
        "anyone you want to skip, then confirm."
    )

    # Candidates are keyed by email_id; the shared panel expects "id".
    for c in candidates:
        c.setdefault("id", c["email_id"])
        c.setdefault("display_name", c.get("company_name", ""))

    chosen = _render_email_review(
        candidates, prefix="fu", selected=selected,
        extra_columns={"Company": "company_name",
                       "Follow-up #": "follow_up_number"},
    )

    # ── Confirm and send ───────────────────────────────────────────────────
    st.markdown("---")
    active = [s for s in st.session_state.senders if s.get("email") and s.get("password")]
    capacity = len(active) * MAX_EMAILS_PER_SENDER

    if not active:
        st.error("❌ Configure at least one sender account before sending.")
    elif len(chosen) > capacity:
        st.warning(
            f"⚠️ {len(chosen)} selected but your {len(active)} sender(s) can "
            f"only send {capacity} this session. The rest would fail — "
            "deselect some, or add more sender accounts."
        )

    c1, c2 = st.columns([2, 1])
    with c1:
        send_clicked = st.button(
            f"📨 Send {len(chosen)} Follow-up(s) Now",
            type="primary",
            disabled=not chosen or not active,
            use_container_width=True,
            key="fu_send",
        )
    with c2:
        if st.button("Cancel", use_container_width=True, key="fu_cancel"):
            st.session_state.followup_candidates = None
            st.session_state.followup_selected = {}
            st.rerun()

    if send_clicked:
        bar  = st.progress(0.0)
        note = st.empty()

        def _progress(done, total, recipient):
            note.caption(f"Sending {done}/{total} — {recipient}")
            bar.progress(done / total)

        summary = send_followups(
            chosen, active,
            stop_event=st.session_state._stop_event,
            progress_cb=_progress,
        )
        note.empty()
        bar.empty()
        st.session_state.followup_result     = summary
        st.session_state.followup_candidates = None
        st.session_state.followup_selected   = {}
        st.rerun()


def _render_scheduled_queue():
    """Show pending scheduled sends, with the option to cancel them."""
    try:
        pending = get_scheduled_sends("pending")
    except Exception:
        return
    if not pending:
        return

    st.markdown("---")
    st.markdown(f"### 🕐 Scheduled ({len(pending)} pending)")
    if not send_worker.worker_alive():
        st.warning("⚠️ The scheduler is not running — restarting it.")
        send_worker.start_worker()

    for job in pending:
        due = job["scheduled_for"]
        overdue = due <= datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        c1, c2 = st.columns([4, 1])
        with c1:
            st.markdown(
                f"**{job['label'] or job['kind']}** — "
                f"{'⏳ due now, sending shortly' if overdue else f'sends {due}'}"
            )
        with c2:
            if st.button("Cancel", key=f"sched_cancel_{job['id']}",
                         use_container_width=True):
                if cancel_scheduled_send(job["id"]):
                    st.success("Scheduled send cancelled.")
                else:
                    st.warning("Too late — that job already started.")
                st.rerun()


def _render_launch_controls(on_start_fn):
    """
    Renders the activity log + START / STOP / Follow-up buttons.
    on_start_fn() is called (with no args) when credentials pass validation.
    """
    st.markdown("---")
    st.markdown("### 📋 Activity Log")

    log_box = st.empty()
    st.progress(0)

    log_box.markdown(
        "<div class='log-box'>" + "\n".join(get_recent_logs(40)) + "</div>",
        unsafe_allow_html=True,
    )

    snap = campaign_runner.snapshot()
    # Checkpoints are stored in SQLite, so they remain available after a
    # Streamlit/browser restart rather than only during this process.
    if snap["status"] == "idle":
        for item in get_paused_campaigns():
            stage = (item.get("resume_stage") or "sending").replace("_", " ")
            a, b = st.columns([4, 1])
            with a:
                st.caption(f"Paused campaign #{item['id']} — {stage}")
            with b:
                if st.button("Resume", key=f"resume_saved_{item['id']}", use_container_width=True):
                    active = [s for s in st.session_state.senders if s.get("email") and s.get("password")]
                    if active and campaign_runner.resume_campaign(
                        item["id"], active, ai_api_key=st.session_state.get("api_key", ""),
                        offering=st.session_state.get("offering", ""),
                        company_name=st.session_state.get("cname", "")):
                        st.rerun()
                    elif not active:
                        st.error("Configure active senders first.")
                    else:
                        st.error("This campaign has no recoverable checkpoint.")
    busy = snap["status"] in ("running", "sending")
    review_pending = st.session_state.campaign_prepared is not None

    c1, c2, c3 = st.columns([2, 1, 1])
    with c1:
        start_clicked = st.button(
            "🚀 START CAMPAIGN",
            disabled=busy or review_pending,
            use_container_width=True,
        )
    with c2:
        stop_clicked = st.button(
            "⏸️ PAUSE", disabled=not busy, use_container_width=True,
        )
    with c3:
        followup_clicked = st.button(
            "🔄 Run Follow-ups", use_container_width=True,
            disabled=review_pending or busy,
        )
    st.caption(
        "Both buttons compose the emails and show them for your approval first "
        "— nothing is sent until you confirm."
    )

    if stop_clicked and busy:
        campaign_runner.request_stop()
        st.warning("Pause requested — halting after the current item.")
        st.rerun()

    if start_clicked and not busy:
        active_senders = [s for s in st.session_state.senders if s.get("email") and s.get("password")]
        if not active_senders:
            st.error("❌ Configure at least one email sender account first.")
        else:
            with st.spinner("Validating sender credentials..."):
                results = [(s, test_sender(s)) for s in active_senders]
            failed = [(s["email"], r["error"]) for s, r in results if not r["ok"]]

            if failed:
                for addr, err in failed:
                    st.error(f"❌ {addr}: {err}")
                st.info("Fix the errors above, then try again.")
            else:
                cs = st.session_state.get("custom_subject", "").strip()
                cb = st.session_state.get("custom_body", "").strip()
                if (cs and not cb) or (cb and not cs):
                    st.error("❌ Fill in BOTH Subject and Body, or leave both empty.")
                else:
                    on_start_fn()
                    st.rerun()

    if followup_clicked:
        # Sync inbox replies first, then fetch candidates for approval
        with st.spinner("Syncing inbox replies & checking who is due a follow-up…"):
            if st.session_state.senders:
                try:
                    fetch_all_replies(st.session_state.senders)
                except Exception as exc:
                    log_warning(f"Inbox reply fetch before follow-ups: {exc}")
            st.session_state.followup_candidates = get_followup_candidates(
                follow_up_after_days=follow_up_days,
                api_key=st.session_state.get("api_key", ""),
                offering=st.session_state.get("offering", ""),
                sender_name=st.session_state.get("cname", ""),
                sender_company=st.session_state.get("cname", ""),
            )
        st.session_state.followup_selected = {
            c["email_id"]: True for c in st.session_state.followup_candidates
        }
        # Fresh widget identities, so a previous preview's edits cannot
        # reappear in this one.
        st.session_state["fu_versions"] = {}
        st.session_state.followup_result = None
        st.rerun()

    cresult = st.session_state.campaign_result
    if cresult:
        st.success(
            f"✅ Campaign sent {cresult['sent']}/{cresult['total']} approved email(s)."
        )
        st.session_state.campaign_result = None

    if snap["status"] == "awaiting_review":
        _adopt_prepared_for_review(snap)
    _render_campaign_review()

    result = st.session_state.followup_result
    if result:
        if result.get("scheduled"):
            st.success(
                f"🕐 Scheduled {result['count']} follow-up(s) for "
                f"**{result['when']}**. Nothing has been sent yet — you can "
                "cancel it below until it runs."
            )
        else:
            st.success(
                f"✅ Follow-ups sent: {result['sent']} | "
                f"failed: {result['failed']} | skipped: {result['skipped']}"
            )
            for err in result.get("errors", [])[:5]:
                st.warning(f"⚠️ {err}")
        st.session_state.followup_result = None

    _render_followup_review()

    # ── Live progress & Pause/Resume state, polled from background thread ──
    if snap["status"] in ("running", "sending"):
        st.info(
            f"⏳ **{snap['step'] or 'Working…'}**  \n"
            f"{snap['detail']}"
            + (f"  \nLeads found: {snap['leads_found']}" if snap["leads_found"] else "")
            + (f"  |  Sent: {snap['sent']}" if snap["status"] == "sending" else "")
        )
        if snap["total"]:
            st.progress(min(1.0, snap["done"] / snap["total"]),
                        text=f"{snap['done']} / {snap['total']}")
        if snap["stopping"]:
            st.warning("Pausing after the current item…")
        st.caption(
            f"Started {snap['started_at']}. This runs in the background — "
            "you can navigate away and come back."
        )
        time.sleep(2)
        st.rerun()

    elif snap["status"] == "paused":
        st.warning(
            f"⏸️ **{snap['step']}**  \n"
            f"{snap['detail']}"
        )
        pc1, pc2 = st.columns(2)
        with pc1:
            if st.button("▶️ Resume Campaign", type="primary", use_container_width=True, key="btn_resume_snap"):
                active_senders = [s for s in st.session_state.senders if s.get("email") and s.get("password")]
                if not active_senders:
                    st.error("❌ Configure active senders first.")
                else:
                    campaign_runner.resume_campaign(
                        snap.get("campaign_id"), active_senders,
                        ai_api_key=st.session_state.get("api_key", ""),
                        offering=st.session_state.get("offering", ""),
                        company_name=st.session_state.get("cname", ""),
                    )
                    st.rerun()
        with pc2:
            if st.button("🗑️ Cancel & Discard", use_container_width=True, key="btn_cancel_snap"):
                if snap.get("campaign_id"):
                    update_campaign_status(snap["campaign_id"], "cancelled")
                campaign_runner.reset()
                st.rerun()

    elif snap["status"] == "error":
        st.error(f"❌ Campaign ended with error: {snap['error']}")
        if st.button("Dismiss", key="camp_dismiss_err"):
            campaign_runner.reset()
            st.rerun()

    elif snap["status"] == "cancelled":
        st.warning("Campaign cancelled. Nothing further was sent.")
        if st.button("Dismiss", key="camp_dismiss_cancel"):
            campaign_runner.reset()
            st.rerun()

    elif snap["status"] == "done":
        st.success(f"✅ {snap['step']}.")
        if st.button("Dismiss", key="camp_dismiss_done"):
            campaign_runner.reset()
            st.rerun()


# ── REPORT BUILDERS ────────────────────────────────────────────────────────────

def _build_styled_excel_report(data: dict) -> bytes:
    """Generate a professionally styled multi-sheet Excel (.xlsx) report using openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    font_title  = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold   = Font(name="Segoe UI", size=10, bold=True)
    font_normal = Font(name="Segoe UI", size=10)

    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_zebra  = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )

    align_center = Alignment(horizontal="center", vertical="center")

    # 1. Executive Summary Sheet
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = f"DAILY OUTBOUND PERFORMANCE REPORT - {data['date']}"
    ws_sum["A1"].font = font_title

    ws_sum.append([])
    ws_sum.append(["Metric", "Count / Value"])
    ws_sum.cell(row=3, column=1).font = font_header
    ws_sum.cell(row=3, column=1).fill = fill_header
    ws_sum.cell(row=3, column=2).font = font_header
    ws_sum.cell(row=3, column=2).fill = fill_header

    metrics = [
        ("Report Date", data["date"]),
        ("Total Emails Sent", data["total_sent"]),
        ("Total Replies Received", data["total_replies"]),
        ("Total Follow-ups Sent", data["total_followups"]),
        ("New Leads Discovered", data["total_leads_scraped"]),
    ]
    for label, val in metrics:
        ws_sum.append([label, val])
        r = ws_sum.max_row
        ws_sum.cell(row=r, column=1).font = font_bold
        ws_sum.cell(row=r, column=2).font = font_normal
        ws_sum.cell(row=r, column=1).border = thin_border
        ws_sum.cell(row=r, column=2).border = thin_border

    ws_sum.append([])
    ws_sum.append(["Sender Account", "Emails Sent Today"])
    sb_row = ws_sum.max_row
    ws_sum.cell(row=sb_row, column=1).font = font_header
    ws_sum.cell(row=sb_row, column=1).fill = fill_header
    ws_sum.cell(row=sb_row, column=2).font = font_header
    ws_sum.cell(row=sb_row, column=2).fill = fill_header

    if data["sender_breakdown"]:
        for sb in data["sender_breakdown"]:
            ws_sum.append([sb["sender_email"], sb["count"]])
            r = ws_sum.max_row
            ws_sum.cell(row=r, column=1).font = font_normal
            ws_sum.cell(row=r, column=2).font = font_bold
            ws_sum.cell(row=r, column=1).border = thin_border
            ws_sum.cell(row=r, column=2).border = thin_border
    else:
        ws_sum.append(["No emails sent on this date", 0])

    for col in ws_sum.columns:
        max_l = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sum.column_dimensions[col_letter].width = max(max_l + 4, 16)

    # 2. Sent Emails Sheet
    ws_sent = wb.create_sheet(title="Sent Emails")
    ws_sent.views.sheetView[0].showGridLines = True
    ws_sent.append(["Time Sent", "Recipient Email", "Sender Email", "Company Name", "Subject", "Status", "Follow-up Sent"])
    for cell in ws_sent[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    for idx, e in enumerate(data["sent_emails"], 2):
        ws_sent.append([
            e.get("sent_at", ""),
            e.get("recipient_email", ""),
            e.get("sender_email", ""),
            e.get("company_name", ""),
            e.get("subject", ""),
            e.get("status", "sent"),
            "Yes" if e.get("follow_up_sent") else "No"
        ])
        for c in range(1, 8):
            cell = ws_sent.cell(row=idx, column=c)
            cell.font = font_normal
            cell.border = thin_border
            if idx % 2 == 1:
                cell.fill = fill_zebra

    for col in ws_sent.columns:
        max_l = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sent.column_dimensions[col_letter].width = min(max(max_l + 3, 12), 45)

    # 3. Replies Received Sheet
    ws_rep = wb.create_sheet(title="Replies Received")
    ws_rep.views.sheetView[0].showGridLines = True
    ws_rep.append(["Received Time", "From Email", "From Name", "Mailbox Account", "Subject", "Snippet"])
    for cell in ws_rep[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    for idx, r in enumerate(data["replies"], 2):
        body_snippet = (r.get("body") or "").replace("\n", " ")[:100]
        ws_rep.append([
            r.get("received_at", ""),
            r.get("from_email", ""),
            r.get("from_name", ""),
            r.get("mailbox", ""),
            r.get("subject", ""),
            body_snippet
        ])
        for c in range(1, 7):
            cell = ws_rep.cell(row=idx, column=c)
            cell.font = font_normal
            cell.border = thin_border
            if idx % 2 == 1:
                cell.fill = fill_zebra

    for col in ws_rep.columns:
        max_l = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_rep.column_dimensions[col_letter].width = min(max(max_l + 3, 12), 50)

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def send_report_email(target_email: str, subject: str, body_text: str, attachment_bytes: bytes, attachment_name: str, senders: list[dict]) -> dict:
    """Send formatted daily report via SMTP with an attached Excel (.xlsx) file."""
    active = [s for s in senders if s.get("email") and s.get("password")]
    if not active:
        return {"success": False, "error": "No configured sender accounts with email & password."}

    sender = active[0]
    email_addr = sanitize_credential(sender.get("email", ""))
    password   = sanitize_credential(sender.get("password", ""))
    server     = sender.get("smtp_server", "")
    port       = int(sender.get("smtp_port", 587))
    use_ssl    = sender.get("use_ssl", port == 465)

    if not server or not email_addr or not password:
        return {"success": False, "error": "Invalid sender account credentials."}

    msg = MIMEMultipart()
    sender_name = sender.get("name", "").strip()
    msg["From"] = f"{sender_name} <{email_addr}>" if sender_name else email_addr
    msg["To"] = target_email
    msg["Subject"] = subject
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid()

    msg.attach(MIMEText(body_text, "plain", "utf-8"))

    part = MIMEApplication(attachment_bytes, Name=attachment_name)
    part["Content-Disposition"] = f'attachment; filename="{attachment_name}"'
    msg.attach(part)

    ctx = ssl.create_default_context()
    try:
        if use_ssl:
            with smtplib.SMTP_SSL(server, port, context=ctx, timeout=20) as s:
                s.login(email_addr, password)
                s.sendmail(email_addr, target_email, msg.as_bytes())
        else:
            with smtplib.SMTP(server, port, timeout=20) as s:
                s.ehlo()
                s.starttls(context=ctx)
                s.ehlo()
                s.login(email_addr, password)
                s.sendmail(email_addr, target_email, msg.as_bytes())
        return {"success": True, "sender_email": email_addr}
    except Exception as exc:
        return {"success": False, "error": str(exc)}


# ══════════════════════════════════════════════════════════════════════════════
# PAGES
# ══════════════════════════════════════════════════════════════════════════════

st.markdown("# AI Sales Machine")
st.markdown("*Automated outbound sales — AI-powered or your own templates*")
st.markdown("---")


# ─────────────────────────────────────────────────────────────────────────────
if page == "Scraper Campaign":

    st.markdown("##  Scraper Campaign")
    st.markdown("Automatically finds businesses on Google Maps, extracts their emails, and sends your campaign.")

    st.markdown("###  Target")
    c1, c2 = st.columns(2)
    with c1:
        business_type = st.text_input("Business Type", key="biz_type",
                                       placeholder="e.g.  restaurant,hotel")
    with c2:
        location = st.text_input("Location", key="loc",
                                  placeholder="e.g. delhi , noida")
    col_scroll, col_skip = st.columns(2)
    with col_scroll:
        max_scrolls = st.number_input("Max Scrolls (more = more leads)", min_value=1, max_value=10000, value=15, step=1, key="max_scrolls")
    with col_skip:
        skip_already_sent = st.checkbox(
            "🚫 Skip previously emailed contacts", value=True,
            help="Exclude prospects who were already sent an email in past campaigns"
        )

    missing = [f for f, v in [("Business Type", business_type), ("Location", location),
                                ("Offering", offering), ("Company Name", company_name)] if not v]
    if missing:
        st.warning(f"⚠️ Complete in sidebar: **{', '.join(missing)}**")

    st.markdown("---")
    render_email_composer(context="scraper")

    def _start_scraper():
        _launch_campaign(
            "scraper",
            business_type=business_type, location=location,
            offering=offering, company_name=company_name,
            company_website=company_website, ai_api_key=ai_api_key,
            max_scrolls=max_scrolls,
            skip_already_sent=skip_already_sent,
            custom_subject=st.session_state.get("custom_subject", "").strip(),
            custom_body=st.session_state.get("custom_body", "").strip(),
        )

    _render_launch_controls(on_start_fn=_start_scraper)


# ─────────────────────────────────────────────────────────────────────────────
elif page == "List Campaign":

    st.markdown("## 📋 List Campaign")
    st.markdown(
        "Upload your own contact list (CSV or Excel). "
        "The system sends your email to every contact, rotating across all your sender accounts."
    )

    # ── CSV FORMAT GUIDE ──────────────────────────────────────────────────────
    st.markdown("### 📂 Upload Your Contact List")

    with st.expander("📖 Required CSV format (click to see)", expanded=False):
        st.markdown("""
Your file must have at least an **email** column. All other columns are optional.

| Column name | Required | Notes |
|---|---|---|
| `email` | ✅ Yes | Must be a valid email address |
| `first_name` | Recommended | Used for personalisation: "Hi John," |
| `last_name` | Optional | Available as `{last_name}` placeholder |
| `company` | Optional | Available as `{company_name}` placeholder |

**Accepted column name variants:**
- Email: `email`, `email_address`, `e-mail`
- First name: `first_name`, `firstname`, `first`, `fname`
- Last name: `last_name`, `lastname`, `last`, `lname`, `surname`
- Company: `company`, `company_name`, `organisation`, `organization`, `business`

**File formats accepted:** `.csv`, `.xlsx`, `.xls`
        """)
        st.download_button(
            label="⬇️ Download sample CSV template",
            data=generate_sample_csv(),
            file_name="contacts_template.csv",
            mime="text/csv",
        )

    # ── FILE UPLOADER ─────────────────────────────────────────────────────────
    uploaded_file = st.file_uploader(
        "Upload CSV or Excel file",
        type=["csv", "xlsx", "xls"],
        key="contact_file",
        help="Columns: email (required), first_name, last_name, company (all optional except email)",
    )

    contacts = st.session_state.get("uploaded_contacts", [])

    if uploaded_file is not None:
        file_bytes = uploaded_file.read()
        parsed = parse_contact_list(file_bytes, uploaded_file.name)

        if parsed:
            parsed, warnings = validate_contact_list(parsed)
            st.session_state.uploaded_contacts = parsed
            contacts = parsed

            for w in warnings:
                st.warning(f"⚠️ {w}")

            st.success(f"✅ {len(contacts)} contacts loaded from **{uploaded_file.name}**")

            # Preview table
            preview_df = pd.DataFrame(contacts[:10])
            st.markdown(f"**Preview (first {min(10, len(contacts))} rows):**")
            st.dataframe(preview_df, use_container_width=True, hide_index=True)
            if len(contacts) > 10:
                st.caption(f"... and {len(contacts) - 10} more rows not shown.")
        else:
            st.error(
                "❌ Could not parse the file. "
                "Make sure it has an 'email' column and is a valid CSV or Excel file."
            )

    elif contacts:
        st.info(f"📋 Using previously uploaded list — **{len(contacts)} contacts** ready.")
        if st.button("🗑️ Clear list"):
            st.session_state.uploaded_contacts = []
            st.rerun()

    else:
        st.markdown(
            "<div class='list-info'>📭 No list uploaded yet. "
            "Upload a CSV or Excel file above to get started.</div>",
            unsafe_allow_html=True,
        )

    # ── SENDER + COMPOSER ────────────────────────────────────────────────────
    if contacts:
        st.markdown("---")
        render_email_composer(context="list")

        # ── SUMMARY METRICS ───────────────────────────────────────────────────
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("👥 Contacts",  len(contacts))
        c2.metric("📤 Senders",   len([s for s in st.session_state.senders if s.get("email")]))
        c3.metric("📅 Follow-up", f"{follow_up_days}d")

        has_names = sum(1 for c in contacts if c.get("first_name"))
        c4.metric("🪪 With First Name", f"{has_names}/{len(contacts)}")

        # ── LAUNCH ────────────────────────────────────────────────────────────
        list_skip_already_sent = st.checkbox(
            "🚫 Skip previously emailed contacts", value=True, key="list_skip_sent",
            help="Exclude contacts from this CSV list who were already sent an email in a past campaign"
        )

        missing_list = [f for f, v in [("Offering", offering), ("Company Name", company_name)] if not v]
        if missing_list:
            st.warning(f"⚠️ Complete in sidebar: **{', '.join(missing_list)}**")

        def _start_list():
            _launch_campaign(
                "list",
                contacts=list(st.session_state.uploaded_contacts),
                offering=offering, company_name=company_name,
                company_website=company_website, ai_api_key=ai_api_key,
                skip_already_sent=list_skip_already_sent,
                custom_subject=st.session_state.get("custom_subject", "").strip(),
                custom_body=st.session_state.get("custom_body", "").strip(),
            )

        # Disable launch if required sidebar fields are missing
        if not missing_list:
            _render_launch_controls(on_start_fn=_start_list)
        else:
            st.info("Fill in the sidebar fields above before launching.")


# ─────────────────────────────────────────────────────────────────────────────
elif page.startswith("Replies"):

    st.markdown("## Replies")
    st.markdown(
        "Outgoing email carries no Reply-To, so prospects reply to whichever "
        "account sent them. This gathers those replies from every mailbox "
        "into one place so you can answer from here."
    )

    senders = [s for s in st.session_state.senders
               if s.get("email") and s.get("password")]

    c1, c2, c3 = st.columns([2, 1.2, 1], vertical_alignment="bottom")
    with c1:
        check_clicked = st.button("🔄 Check all mailboxes",
                                  type="primary", use_container_width=True,
                                  disabled=not senders)
    with c2:
        days_back = st.number_input("Look back (days)", 1, 180, 30, key="rep_days")
    with c3:
        show_all = st.checkbox("Include handled", value=False, key="rep_all")

    if not senders:
        st.warning("Configure at least one sender account first — replies are "
                   "read from those same mailboxes over IMAP.")

    with st.expander("ℹ️ How the mailbox is reached", expanded=False):
        st.caption(
            "IMAP host is guessed by swapping `smtp.` for `imap.`. Gmail also "
            "requires IMAP enabled in settings plus an app password."
        )
        st.dataframe(pd.DataFrame([
            {"Account": s["email"], "SMTP": s.get("smtp_server", ""),
             "IMAP (used)": imap_host_for(s)} for s in senders
        ]), use_container_width=True, hide_index=True)

    if check_clicked:
        bar, note = st.progress(0.0), st.empty()

        def _progress(mailbox, n, total):
            note.caption(f"{mailbox} — message {n}/{total}")
            if total:
                bar.progress(min(1.0, n / total))

        with st.spinner("Connecting to mailboxes…"):
            results = fetch_all_replies(senders, days_back=int(days_back),
                                        progress_cb=_progress)
        note.empty(); bar.empty()

        total_new = sum(r["new"] for r in results)
        if total_new:
            st.success(f"✅ Found {total_new} new reply(ies).")
        else:
            st.info("No new replies found.")
        for r in results:
            if r["error"]:
                st.error(f"❌ {r['mailbox']}: {r['error']}")
            else:
                st.caption(
                    f"{r['mailbox']} — scanned {r['scanned']}, "
                    f"new {r['new']}, auto-replies skipped {r['skipped_auto']}"
                )
        st.rerun()

    st.markdown("---")
    replies = get_replies(only_unhandled=not show_all)

    if not replies:
        st.info("No replies stored yet. Run a check above.")
    else:
        unread = sum(1 for r in replies if not r["is_read"])
        total_count = len(replies)
        reply_label = "1 Reply" if total_count == 1 else f"{total_count} Replies"

        m_col1, m_col2 = st.columns(2)
        with m_col1:
            st.markdown(
                f"<div class='metric-card'><div class='metric-value'>📬 {total_count}</div>"
                f"<div class='metric-label'>{reply_label}</div></div>",
                unsafe_allow_html=True,
            )
        with m_col2:
            st.markdown(
                f"<div class='metric-card'><div class='metric-value'>🔵 {unread}</div>"
                f"<div class='metric-label'>Unread Replies</div></div>",
                unsafe_allow_html=True,
            )
        st.markdown("<br>", unsafe_allow_html=True)

        labels = [
            f"{'🔵 ' if not r['is_read'] else '   '}"
            f"{r['from_email']} — {(r['subject'] or '(no subject)')[:46]}"
            f"  · {(r['received_at'] or '')[:16]}"
            for r in replies
        ]
        pick_idx = st.selectbox(
            "Select a reply",
            range(len(replies)),
            index=None,
            placeholder="Select a reply",
            format_func=lambda i: labels[i],
            key="rep_pick",
        )

        if pick_idx is not None:
            rep = replies[pick_idx]

            if not rep["is_read"]:
                set_reply_flags(rep["id"], is_read=True)

            meta_c1, meta_c2 = st.columns(2)
            with meta_c1:
                st.markdown(
                    f"**From:** {rep['from_name'] or ''} &lt;{rep['from_email']}&gt;  \n"
                    f"**Arrived in:** `{rep['mailbox']}`",
                    unsafe_allow_html=True,
                )
            with meta_c2:
                st.markdown(
                    f"**Received:** {rep['received_at']}  \n"
                    f"**Subject:** {rep['subject']}",
                    unsafe_allow_html=True,
                )
            st.text_area("Their message", value=rep["body"] or "(empty)",
                         height=220, disabled=True, key=f"rep_body_{rep['id']}")

            # ── Reply to them ──────────────────────────────────────────────────
            st.markdown("#### ✍️ Reply")
            mailbox = next((s for s in senders
                            if s["email"].lower() == (rep["mailbox"] or "").lower()), None)
            if not mailbox:
                st.warning(
                    f"The mailbox `{rep['mailbox']}` is not in your configured "
                    "senders, so this reply cannot be sent from the address the "
                    "prospect knows. Add that account in Manage Senders."
                )
            else:
                subj_default = rep["subject"] or ""
                if not subj_default.lower().startswith("re:"):
                    subj_default = f"Re: {subj_default}"

                r_subject = st.text_input("Subject", value=subj_default,
                                          key=f"rep_subj_{rep['id']}")
                r_body = st.text_area("Your reply", height=200,
                                      key=f"rep_reply_{rep['id']}",
                                      placeholder="Type your response…")
                st.caption(f"Sends from **{mailbox['email']}** — the same address "
                           "they replied to. Your signature is appended.")

                b1, b2 = st.columns([2, 1])
                with b1:
                    send_reply = st.button("📨 Send reply", type="primary",
                                           use_container_width=True,
                                           disabled=not r_body.strip(),
                                           key=f"rep_send_{rep['id']}")
                with b2:
                    if st.button("✅ Mark handled", use_container_width=True,
                                 key=f"rep_done_{rep['id']}"):
                        set_reply_flags(rep["id"], handled=True)
                        st.rerun()

                if send_reply:
                    with st.spinner("Sending…"):
                        result = SMTPSender([mailbox]).send_direct(
                            sender=mailbox, to_email=rep["from_email"],
                            subject=r_subject, body=r_body,
                            in_reply_to=rep["message_id"] or "",
                        )
                    if result["success"]:
                        insert_email_sent(
                            lead_id=rep.get("lead_id"), company_name="",
                            recipient_email=rep["from_email"],
                            sender_email=mailbox["email"],
                            subject=r_subject, body=r_body,
                            message_id=result.get("message_id"),
                        )
                        set_reply_flags(rep["id"], handled=True)
                        st.success(f"✅ Reply sent to {rep['from_email']}.")
                        st.rerun()
                    else:
                        st.error(f"❌ {result.get('error')}")


# ─────────────────────────────────────────────────────────────────────────────
elif page == "Dashboard":

    st.markdown("## 📊 Dashboard & Reports")
    tab1, tab2 = st.tabs(["📊 Overall Analytics", "📅 Daily Report Generator"])

    with tab1:
        stats = get_campaign_stats()

        cols  = st.columns(5)
        cards = [
            ("Total Leads",  stats["total_leads"]),
            ("Emails Sent",  stats["sent"]),
            ("Opened",       f"{stats['opened']} ({stats['open_rate']}%)"),
            ("Replied",      f"{stats['replied']} ({stats['reply_rate']}%)"),
            ("Follow-ups",   stats["follow_ups"]),
        ]
        for col, (label, value) in zip(cols, cards):
            with col:
                st.markdown(
                    f"<div class='metric-card'><div class='metric-value'>{value}</div>"
                    f"<div class='metric-label'>{label}</div></div>",
                    unsafe_allow_html=True,
                )

        st.markdown("<br>", unsafe_allow_html=True)

        if stats["sent"] > 0:
            c1, c2 = st.columns(2)
            with c1:
                fig = go.Figure(go.Pie(
                    labels=["Sent","Opened","Replied","Follow-ups"],
                    values=[stats["sent"],stats["opened"],stats["replied"],stats["follow_ups"]],
                    hole=0.4, marker_colors=["#0099ff","#00d4aa","#ff9f43","#ee5a24"],
                ))
                fig.update_layout(title="Status Breakdown", paper_bgcolor="rgba(0,0,0,0)",
                                  font_color="white", height=350)
                st.plotly_chart(fig, use_container_width=True)
            with c2:
                fig2 = go.Figure(go.Bar(
                    x=["Sent","Opened","Replied","Follow-ups"],
                    y=[stats["sent"],stats["opened"],stats["replied"],stats["follow_ups"]],
                    marker_color=["#0099ff","#00d4aa","#ff9f43","#ee5a24"],
                ))
                fig2.update_layout(title="Counts", paper_bgcolor="rgba(0,0,0,0)",
                                   plot_bgcolor="rgba(0,0,0,0)", font_color="white",
                                   height=350, yaxis=dict(gridcolor="#333"))
                st.plotly_chart(fig2, use_container_width=True)

        st.markdown("### 📋 Sent Emails")
        emails = get_all_emails_sent()
        if emails:
            df = pd.DataFrame(emails)
            df.columns = [c.replace("_"," ").title() for c in df.columns]
            df["Follow Up Sent"] = df["Follow Up Sent"].map({1:"✅ Yes", 0:"—"})
            st.dataframe(df, use_container_width=True, height=420, hide_index=True)
        else:
            st.info("No emails sent yet. Launch a campaign to get started.")

        db_btn_col1, db_btn_col2 = st.columns([1, 2])
        with db_btn_col1:
            if st.button("🔄 Refresh Data", use_container_width=True):
                st.rerun()
        with db_btn_col2:
            if st.button("🔄 Run Follow-ups", type="primary", use_container_width=True, key="dash_run_fu_btn"):
                with st.spinner("Syncing inbox replies & checking who is due a follow-up…"):
                    if st.session_state.senders:
                        try:
                            fetch_all_replies(st.session_state.senders)
                        except Exception as exc:
                            log_warning(f"Inbox reply fetch before follow-ups: {exc}")
                    st.session_state.followup_candidates = get_followup_candidates(
                        follow_up_after_days=follow_up_days,
                        api_key=st.session_state.get("api_key", ""),
                        offering=st.session_state.get("offering", ""),
                        sender_name=st.session_state.get("cname", ""),
                        sender_company=st.session_state.get("cname", ""),
                    )
                st.session_state.followup_selected = {
                    c["email_id"]: True for c in st.session_state.followup_candidates
                }
                st.session_state["fu_versions"] = {}
                st.session_state.followup_result = None
                st.rerun()

        result = st.session_state.followup_result
        if result:
            if result.get("scheduled"):
                st.success(
                    f"🕐 Scheduled {result['count']} follow-up(s) for "
                    f"**{result['when']}**. Nothing has been sent yet."
                )
            else:
                st.success(
                    f"✅ Follow-ups sent: {result['sent']} | "
                    f"failed: {result['failed']} | skipped: {result['skipped']}"
                )
                for err in result.get("errors", [])[:5]:
                    st.warning(f"⚠️ {err}")
            st.session_state.followup_result = None

        _render_followup_review()

    with tab2:
        st.markdown("### 📅 Daily Activity & Performance Report")
        st.markdown("Select any date to view and download a comprehensive 1-click report of emails sent, replies received, follow-ups executed, and new leads discovered.")

        r_col1, r_col2 = st.columns([2, 1])
        with r_col1:
            report_date = st.date_input("Select Date for Report", value=date.today(), key="daily_rep_date")
        with r_col2:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("🔄 Load Report", type="primary", use_container_width=True):
                st.rerun()

        target_date_str = report_date.strftime("%Y-%m-%d")
        daily_data = get_daily_report_data(target_date_str)

        # Summary Metrics Cards
        mc1, mc2, mc3, mc4 = st.columns(4)
        mc1.metric("📤 Sent Today", daily_data["total_sent"])
        mc2.metric("📬 Replies Received", daily_data["total_replies"])
        mc3.metric("🔄 Follow-ups Sent", daily_data["total_followups"])
        mc4.metric("🔍 Leads Discovered", daily_data["total_leads_scraped"])

        st.markdown("---")
        st.markdown("#### 📥 Download Report File")

        excel_data = _build_styled_excel_report(daily_data)
        st.download_button(
            label=f"📊 Download Excel Report (.xlsx)",
            data=excel_data,
            file_name=f"daily_report_{target_date_str}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            type="primary",
        )

        st.markdown("---")
        st.markdown("#### 📧 Email Daily Report")
        st.markdown("Send this report (with the styled Excel attachment `.xlsx`) directly to any email address.")
        ec1, ec2 = st.columns([3, 1])
        with ec1:
            target_report_email = st.text_input(
                "Recipient Email Address",
                placeholder="Enter recipient email (e.g. boss@company.com)",
                key="target_report_email_input",
                label_visibility="collapsed"
            )
        with ec2:
            if st.button("📨 Send Report", type="primary", use_container_width=True):
                if not target_report_email or "@" not in target_report_email:
                    st.error("Please enter a valid recipient email address.")
                else:
                    with st.spinner("Sending report email with Excel attachment..."):
                        sub = f"📊 Daily Outbound Performance Report - {target_date_str}"
                        body_txt = (
                            f"Hello,\n\n"
                            f"Please find attached the Daily Outbound Performance Report for {target_date_str}.\n\n"
                            f"Executive Summary:\n"
                            f"  - Total Emails Sent     : {daily_data['total_sent']}\n"
                            f"  - Total Replies Received: {daily_data['total_replies']}\n"
                            f"  - Total Follow-ups Sent : {daily_data['total_followups']}\n"
                            f"  - Leads Discovered      : {daily_data['total_leads_scraped']}\n\n"
                            f"Detailed breakdown sheets (Summary, Sent Emails, Replies Received) are attached in the Excel file.\n\n"
                            f"Best regards,\n"
                            f"AI Sales Machine"
                        )
                        res = send_report_email(
                            target_email=target_report_email.strip(),
                            subject=sub,
                            body_text=body_txt,
                            attachment_bytes=excel_data,
                            attachment_name=f"daily_report_{target_date_str}.xlsx",
                            senders=st.session_state.senders
                        )
                        if res.get("success"):
                            st.success(f"✅ Report successfully sent to **{target_report_email.strip()}** via `{res.get('sender_email')}`!")
                        else:
                            st.error(f"❌ Failed to send report: {res.get('error')}")

        st.markdown("---")

        # Sender Volume Breakdown Table
        if daily_data["sender_breakdown"]:
            st.markdown("#### 📧 Sender Account Volume Breakdown")
            df_senders = pd.DataFrame(daily_data["sender_breakdown"])
            df_senders.columns = ["Sender Account", "Emails Sent"]
            st.dataframe(df_senders, use_container_width=True, hide_index=True)

        # Sent Emails Breakdown Table
        st.markdown(f"#### 📤 Sent Emails on {target_date_str} ({len(daily_data['sent_emails'])})")
        if daily_data["sent_emails"]:
            df_sent = pd.DataFrame(daily_data["sent_emails"])
            df_sent.columns = [c.replace("_", " ").title() for c in df_sent.columns]
            df_sent["Follow Up Sent"] = df_sent["Follow Up Sent"].map({1: "✅ Yes", 0: "—"})
            st.dataframe(df_sent, use_container_width=True, hide_index=True)
        else:
            st.info(f"No emails were sent on {target_date_str}.")

        # Replies Breakdown Table
        st.markdown(f"#### 📬 Replies Received on {target_date_str} ({len(daily_data['replies'])})")
        if daily_data["replies"]:
            df_rep = pd.DataFrame(daily_data["replies"])
            df_rep.columns = [c.replace("_", " ").title() for c in df_rep.columns]
            st.dataframe(df_rep, use_container_width=True, hide_index=True)
        else:
            st.info(f"No replies were recorded on {target_date_str}.")


# ─────────────────────────────────────────────────────────────────────────────
elif page == "Manage Senders":

    render_sender_form(allow_bulk=True)

    st.markdown("---")
    st.markdown("### 📖 SMTP Quick Reference")
    ref = {
        "Provider":    ["Gmail","Hostinger","cPanel Webmail","Outlook","Office 365",
                        "Yahoo","Zoho","SendGrid","Mailgun","Amazon SES","Titan","iCloud"],
        "SMTP Server": ["smtp.gmail.com","smtp.hostinger.com","mail.yourdomain.com",
                        "smtp.live.com","smtp.office365.com","smtp.mail.yahoo.com",
                        "smtp.zoho.com","smtp.sendgrid.net","smtp.mailgun.org",
                        "email-smtp.us-east-1.amazonaws.com","smtp.titan.email","smtp.mail.me.com"],
        # Kept as strings: a mixed int/"N/A" column cannot be serialised to
        # Arrow and renders as an error instead of a table.
        "TLS Port":    ["587"]*12,
        "SSL Port":    ["465","465","465","N/A","N/A","465",
                        "465","465","465","465","465","587"],
        "Auth":        ["App Password","Mailbox PW","Mailbox PW","Regular PW","Regular PW",
                        "App Password","Regular PW","API Key","SMTP PW","SES Credential",
                        "Mailbox PW","App-Specific PW"],
    }
    st.dataframe(pd.DataFrame(ref), use_container_width=True, hide_index=True)

    st.markdown("---")
    st.markdown("### 🔐 App Password Setup Guides")
    with st.expander("Gmail"):
        st.markdown("""
1. Go to **myaccount.google.com** → **Security**
2. Enable **2-Step Verification**
3. Search **App Passwords** → App = Mail, Device = Windows Computer → **Generate**
4. Copy the 16-character password into the Password field
        """)
    with st.expander("Hostinger / cPanel"):
        st.markdown("""
1. hPanel or cPanel → **Email** → **Email Accounts**
2. Full email address as username, the mailbox password
3. SMTP: `smtp.hostinger.com` port 587 (TLS) or 465 (SSL)
        """)
    with st.expander("Yahoo / Zoho / iCloud"):
        st.markdown("""
- **Yahoo**: myaccount.yahoo.com → Security → Generate app password
- **Zoho**: Regular password works; use `smtp.zoho.in` if on Indian servers
- **iCloud**: appleid.apple.com → App-Specific Passwords
        """)
    with st.expander("SendGrid / Mailgun / Amazon SES"):
        st.markdown("""
- **SendGrid**: Username = `apikey`, Password = your API key
- **Mailgun**: Use SMTP credentials from your Mailgun dashboard
- **SES**: AWS Console → SES → SMTP Settings → Create SMTP Credentials
        """)
